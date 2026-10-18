"""Spotcheck: sample parsed holdings against the raw file they came from.

For the latest included file of every issuer (or one issuer), ``spotcheck``
re-parses the raw file, picks N holdings deterministically from a seed, and
prints the raw weight as written in the file (the CSV cell, the XLSX cell, the
JSON field, or the PDF line text) next to the normalized fraction the build
stored. It also prints the file's weight sum, and for PDFs the disclosure's
stated total against the parsed total. The tool only reads: it never writes to
disk, and the database is opened read-only.

The selection is deterministic for a given ``seed``: the random source is
seeded with the seed, the source id, the fund ticker, the as-of date and the
file's SHA-256, so the same build and seed always produce the same sample.
"""

from __future__ import annotations

import csv
import io
import json
import os
import random
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import duckdb

from etfh.build import connect
from etfh.parse.registry import find_parser
from etfh.registry import Source, load_sources
from etfh.storage import RawStore
from etfh.types import ParsedFile, RawFile
from etfh.units import parse_number

__all__ = ["spotcheck"]

_LATEST_FILES_SQL = """
SELECT fv.source_id, fv.fund_ticker, fv.as_of_date, fv.file_version,
       fv.raw_file_sha256, fv.path, s.issuer
FROM file_versions fv
JOIN sources s ON s.source_id = fv.source_id
WHERE fv.status = 'included'{issuer_filter}
QUALIFY row_number() OVER (
    PARTITION BY s.issuer
    ORDER BY fv.as_of_date DESC, fv.file_version DESC, fv.fund_ticker ASC
) = 1
ORDER BY s.issuer, fv.fund_ticker
"""

_NAME_KEYS = ("name", "security_name", "security", "holding", "description")
_WEIGHT_KEYS = (
    "weight",
    "portfolio_mweight",
    "market_weight",
    "allocation",
    "percentage",
    "pct",
)


@dataclass(frozen=True)
class _Target:
    source_id: str
    fund_ticker: str
    as_of_date: date
    file_version: int
    sha256: str
    path: str
    issuer: str


def spotcheck(
    raw_root: Path,
    db_path: Path,
    *,
    n: int = 5,
    seed: int = 0,
    issuer: str | None = None,
    config_dir: Path | None = None,
) -> str:
    """Return a deterministic spotcheck report for the latest file per issuer.

    ``raw_root`` is the raw store the database was built from; the file is
    re-read from there, so the report reflects exactly the stored bytes. Raises
    ``ValueError`` when ``n`` is not positive.
    """
    if n < 1:
        raise ValueError("n must be at least 1")
    sources = _load_sources(config_dir)
    raws = _raws_by_sha(RawStore(raw_root).iter_raw())
    connection = connect(db_path, read_only=True)
    try:
        targets = _targets(connection, issuer)
    finally:
        connection.close()
    lines = [f"spotcheck n={n} seed={seed}"]
    if not targets:
        lines.append("no included files found" + (f" for issuer {issuer!r}" if issuer else ""))
        return "\n".join(lines) + "\n"
    for target in targets:
        lines.extend(_report_target(target, raws, sources, n=n, seed=seed))
    return "\n".join(lines) + "\n"


def _targets(connection: duckdb.DuckDBPyConnection, issuer: str | None) -> list[_Target]:
    issuer_filter = ""
    parameters: list[Any] = []
    if issuer is not None:
        issuer_filter = " AND s.issuer = ?"
        parameters.append(issuer)
    frame = connection.execute(
        _LATEST_FILES_SQL.format(issuer_filter=issuer_filter), parameters
    ).pl()
    return [
        _Target(
            source_id=str(row["source_id"]),
            fund_ticker=str(row["fund_ticker"]),
            as_of_date=row["as_of_date"],
            file_version=int(row["file_version"]),
            sha256=str(row["raw_file_sha256"]),
            path=str(row["path"]),
            issuer=str(row["issuer"]),
        )
        for row in frame.iter_rows(named=True)
    ]


def _report_target(
    target: _Target,
    raws: dict[str, list[RawFile]],
    sources: dict[str, Source],
    *,
    n: int,
    seed: int,
) -> list[str]:
    raw = _pick_raw(target, raws)
    if raw is None:
        return [
            f"source={target.source_id} issuer={target.issuer} fund={target.fund_ticker} "
            f"as_of={target.as_of_date.isoformat()} error=raw file {target.sha256[:12]} "
            "is not in --raw"
        ]
    source = sources.get(target.source_id)
    parser = find_parser(raw, source) if source is not None else None
    if parser is None:
        return [
            f"source={target.source_id} issuer={target.issuer} fund={target.fund_ticker} "
            f"as_of={target.as_of_date.isoformat()} error=no parser for the stored file"
        ]
    parsed = parser.parse(raw)
    frame = parsed.frame
    scale = _declared_scale(parsed)
    raw_values = _raw_values(raw, parsed, scale=scale)
    weight_sum = float(frame["weight"].sum()) if frame.height else 0.0
    header = (
        f"source={target.source_id} issuer={target.issuer} fund={target.fund_ticker} "
        f"as_of={parsed.as_of_date.isoformat()} version={target.file_version} "
        f"rows={frame.height} weight_sum={weight_sum:.4f} file={raw.path}"
    )
    stated_total = _note_value(parsed.notes, "stated_total_fraction")
    if stated_total is not None:
        header += f" stated_total={stated_total} parsed_total={weight_sum:.4f}"
    lines = [header]
    indices = _sample_indices(frame.height, n=n, seed=seed, target=target)
    for index in indices:
        name = str(frame["security_name"][index])
        weight = float(frame["weight"][index])
        raw_value = raw_values[index] if index < len(raw_values) else None
        rendered = raw_value if raw_value is not None else "(not located)"
        lines.append(f"  {name} | raw={rendered} | weight={weight:.6f}")
    return lines


def _sample_indices(height: int, *, n: int, seed: int, target: _Target) -> list[int]:
    if height == 0:
        return []
    rng = random.Random(
        f"{seed}|{target.source_id}|{target.fund_ticker}|{target.as_of_date.isoformat()}"
        f"|{target.sha256}"
    )
    return sorted(rng.sample(range(height), min(n, height)))


def _raw_values(raw: RawFile, parsed: ParsedFile, *, scale: float) -> list[str | None]:
    extension = raw.ext.lower()
    if extension == "csv":
        return _csv_values(raw.path, parsed, scale=scale)
    if extension == "xlsx":
        return _xlsx_values(raw.path, parsed, scale=scale)
    if extension == "json":
        return _json_values(raw.path, parsed, scale=scale)
    if extension == "pdf":
        return _pdf_values(raw.path, parsed)
    return [None] * parsed.frame.height


def _csv_values(path: Path, parsed: ParsedFile, *, scale: float) -> list[str | None]:
    text = _read_text(path)
    rows = [list(row) for row in csv.reader(io.StringIO(text))]
    return _match_rows(rows, parsed, scale=scale)


def _xlsx_values(path: Path, parsed: ParsedFile, *, scale: float) -> list[str | None]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = [[cell for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    return _match_rows(rows, parsed, scale=scale)


def _match_rows(rows: list[list[Any]], parsed: ParsedFile, *, scale: float) -> list[str | None]:
    """Find, for each parsed row, the file row with its name and closest weight."""
    frame = parsed.frame
    names = [str(name) for name in frame["security_name"].to_list()]
    weights = [float(weight) for weight in frame["weight"].to_list()]
    used: set[int] = set()
    values: list[str | None] = []
    for name, weight in zip(names, weights, strict=True):
        target = weight * scale
        best_index: int | None = None
        best_cell: str | None = None
        best_delta: float | None = None
        for index, row in enumerate(rows):
            if index in used:
                continue
            if not any(name == _text(cell) for cell in row):
                continue
            cell, delta = _closest_cell(row, target)
            if cell is None:
                continue
            if best_delta is None or delta < best_delta:
                best_index, best_cell, best_delta = index, cell, delta
                if delta <= 1e-9:
                    break
        if best_index is not None:
            used.add(best_index)
            values.append(best_cell)
        else:
            values.append(None)
    return values


def _closest_cell(row: list[Any], target: float) -> tuple[str | None, float]:
    best: str | None = None
    best_delta: float | None = None
    for cell in row:
        text = _text(cell)
        if not text:
            continue
        number = parse_number(text)
        if number is None:
            continue
        delta = abs(number - target)
        if best_delta is None or delta < best_delta:
            best, best_delta = text, delta
    if best is None or best_delta is None:
        return None, float("inf")
    return best, best_delta


def _json_values(path: Path, parsed: ParsedFile, *, scale: float) -> list[str | None]:
    payload: Any = json.loads(path.read_text(encoding="utf-8"))
    entries = list(_json_entries(payload))
    frame = parsed.frame
    names = [str(name) for name in frame["security_name"].to_list()]
    weights = [float(weight) for weight in frame["weight"].to_list()]
    used: set[int] = set()
    values: list[str | None] = []
    for name, weight in zip(names, weights, strict=True):
        target = weight * scale
        best_index: int | None = None
        best_value: str | None = None
        best_delta: float | None = None
        for index, entry in enumerate(entries):
            if index in used:
                continue
            if _entry_name(entry) != name:
                continue
            rendered, delta = _entry_weight(entry, target)
            if rendered is None:
                continue
            if best_delta is None or delta < best_delta:
                best_index, best_value, best_delta = index, rendered, delta
                if delta <= 1e-9:
                    break
        if best_index is not None:
            used.add(best_index)
            values.append(best_value)
        else:
            values.append(None)
    return values


def _json_entries(node: Any) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    if isinstance(node, dict):
        if _entry_name(node) is not None:
            entries.append(node)
        for value in node.values():
            entries.extend(_json_entries(value))
    elif isinstance(node, list):
        for item in node:
            entries.extend(_json_entries(item))
    return entries


def _entry_name(entry: dict[str, Any]) -> str | None:
    for key in _NAME_KEYS:
        value = entry.get(key)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def _entry_weight(entry: dict[str, Any], target: float) -> tuple[str | None, float]:
    best: str | None = None
    best_delta: float | None = None
    for key in _WEIGHT_KEYS:
        value = entry.get(key)
        if value is None or isinstance(value, (dict, list)):
            continue
        rendered = str(value).strip()
        number = parse_number(rendered)
        if not rendered or number is None:
            continue
        delta = abs(number - target)
        if best_delta is None or delta < best_delta:
            best, best_delta = rendered, delta
    if best is None or best_delta is None:
        return None, float("inf")
    return best, best_delta


def _pdf_values(path: Path, parsed: ParsedFile) -> list[str | None]:
    import pdfplumber

    lines: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            lines.extend(line.strip() for line in text.splitlines() if line.strip())
    names = [str(name) for name in parsed.frame["security_name"].to_list()]
    used: set[int] = set()
    values: list[str | None] = []
    for name in names:
        match: str | None = None
        for index, line in enumerate(lines):
            if index in used:
                continue
            if name in line:
                match = line
                used.add(index)
                break
        values.append(match)
    return values


def _declared_scale(parsed: ParsedFile) -> float:
    for note in parsed.notes:
        if note == "declared_unit=percent":
            return 100.0
        if note == "declared_unit=fraction":
            return 1.0
    return 1.0


def _note_value(notes: list[str], key: str) -> str | None:
    prefix = f"{key}="
    for note in notes:
        if note.startswith(prefix):
            return note[len(prefix) :]
    return None


def _pick_raw(target: _Target, raws: dict[str, list[RawFile]]) -> RawFile | None:
    candidates = raws.get(target.sha256, [])
    for raw in candidates:
        if str(raw.path) == target.path:
            return raw
    return candidates[0] if candidates else None


def _raws_by_sha(raws: list[RawFile]) -> dict[str, list[RawFile]]:
    by_sha: dict[str, list[RawFile]] = {}
    for raw in raws:
        by_sha.setdefault(raw.sha256, []).append(raw)
    return by_sha


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _load_sources(config_dir: Path | None) -> dict[str, Source]:
    if config_dir is None:
        config_dir = Path(os.environ.get("ETFH_CONFIG_DIR", "config"))
    return load_sources(config_dir / "sources.yaml")
