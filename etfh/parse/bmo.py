"""Parser for BMO Global Asset Management daily holdings XLSX.

The workbook has a single sheet whose first row is a header. The header block
varies by fund type:

- equity funds: ``Weight (%), Name, ISIN, Currency, Number of Shares,
  Base Market Value, Sector, Asset Class``;
- fixed-income funds: ``Weight (%), Name, ISIN, Par Value, Currency, Coupon,
  Maturity Date, Base Market Value, Sector, Asset Class``.

Columns are therefore mapped by name, not position. ``Weight (%)`` is a
percent; cash rows (``Asset Class = Cash``) can carry small negative weights.
The workbook states no portfolio date; the collector records the date in the
file name (the holdings as-of date) as the raw file's as-of hint and the
parser refuses to guess when it is missing.

Values are read with ``polars.read_excel(engine="calamine")`` without a header
so any leading rows survive; openpyxl is the fallback reader.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import polars as pl

from etfh.types import FileFormat, ParsedFile, RawFile
from etfh.units import WeightUnit, check_unit, normalize_weights, parse_number

__all__ = ["PARSERS", "BmoParser"]

_SOURCE_ID = "bmo_daily"
_COLUMNS = [
    "fund_ticker",
    "as_of_date",
    "security_name",
    "security_ticker",
    "exchange",
    "isin",
    "cusip",
    "asset_class",
    "sector",
    "country",
    "currency",
    "weight",
    "shares",
    "market_value",
    "source_id",
    "raw_file_sha256",
]

_ASSET_CLASSES = {
    "equity": "equity",
    "fixed income": "fixed_income",
    "fixed_income": "fixed_income",
    "cash": "cash",
    "derivative": "derivative",
    "derivatives": "derivative",
    "fx": "fx",
    "currency": "fx",
    "foreign currency": "fx",
    "commodity": "commodity",
    "balanced": "balanced",
    "other": "other",
}
_NULL_LABELS = {"", "-", "\u2013", "n/a"}


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _xlsx_rows(path: Path) -> list[list[Any]]:
    """Read every cell of the first sheet, with an openpyxl fallback."""
    try:
        frame = pl.read_excel(
            path,
            engine="calamine",
            has_header=False,
            infer_schema_length=None,
        )
        return [list(row) for row in frame.rows()]
    except Exception:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]


def _header_index(rows: list[list[Any]]) -> int:
    for index, cells in enumerate(rows[:20]):
        labels = {(_text(cell) or "").lower() for cell in cells}
        if "weight (%)" in labels and "name" in labels:
            return index
    raise ValueError("no 'Weight (%)'/'Name' header row found")


def _asset_class(label: str | None, name: str, fixed_income_shape: bool) -> str:
    text = (label or "").strip().lower()
    if text in _ASSET_CLASSES:
        return _ASSET_CLASSES[text]
    if "cash" in name.lower():
        return "cash"
    if fixed_income_shape:
        return "fixed_income"
    return "other"


class BmoParser:
    """Parse a BMO daily holdings workbook into fraction weights."""

    source_id = _SOURCE_ID
    formats = frozenset({FileFormat.xlsx})

    def can_parse(self, raw: RawFile) -> bool:
        return raw.source_id == self.source_id and raw.ext.lower() == "xlsx"

    def parse(self, raw: RawFile) -> ParsedFile:
        as_of_date = raw.as_of_hint
        if as_of_date is None:
            raise ValueError(
                f"{raw.path}: no as-of date; the collector records the date in the file name"
            )
        rows = _xlsx_rows(Path(raw.path))
        header_index = _header_index(rows)
        headers = [(_text(cell) or "") for cell in rows[header_index]]
        positions = {label: index for index, label in enumerate(headers) if label}
        fixed_income_shape = "Par Value" in positions or "Coupon" in positions
        records: list[dict[str, Any]] = []
        for cells in rows[header_index + 1 :]:
            record = {
                label: (cells[index] if index < len(cells) else None)
                for label, index in positions.items()
            }
            name = _text(record.get("Name"))
            if name is None:
                continue
            weight = parse_number(_text(record.get("Weight (%)")))
            if weight is None:
                continue
            sector = _text(record.get("Sector"))
            if sector is not None and sector.lower() in _NULL_LABELS:
                sector = None
            records.append(
                {
                    "security_name": name,
                    "security_ticker": None,
                    "exchange": None,
                    "isin": _text(record.get("ISIN")),
                    "cusip": _text(record.get("CUSIP")),
                    "asset_class": _asset_class(
                        _text(record.get("Asset Class")), name, fixed_income_shape
                    ),
                    "sector": sector or "Unknown",
                    "country": "Unknown",
                    "currency": _text(record.get("Currency")),
                    "weight": weight,
                    "shares": parse_number(_text(record.get("Number of Shares"))),
                    "market_value": parse_number(_text(record.get("Base Market Value"))),
                }
            )
        if not records:
            raise ValueError(f"{raw.path}: no holdings rows found")
        declared = WeightUnit.percent
        weights = [row["weight"] for row in records]
        unit_check = check_unit(pl.Series("weight", weights, dtype=pl.Float64), declared)
        notes = [f"declared_unit={declared.value}"]
        if not unit_check.consistent:
            notes.append(f"unit_check={unit_check.detail()}")
        count = len(records)
        frame = pl.DataFrame(
            {
                "fund_ticker": pl.Series([raw.fund_ticker] * count, dtype=pl.String),
                "as_of_date": pl.Series([as_of_date] * count, dtype=pl.Date),
                "security_name": pl.Series(
                    [row["security_name"] for row in records], dtype=pl.String
                ),
                "security_ticker": pl.Series(
                    [row["security_ticker"] for row in records], dtype=pl.String
                ),
                "exchange": pl.Series([row["exchange"] for row in records], dtype=pl.String),
                "isin": pl.Series([row["isin"] for row in records], dtype=pl.String),
                "cusip": pl.Series([row["cusip"] for row in records], dtype=pl.String),
                "asset_class": pl.Series([row["asset_class"] for row in records], dtype=pl.String),
                "sector": pl.Series([row["sector"] for row in records], dtype=pl.String),
                "country": pl.Series([row["country"] for row in records], dtype=pl.String),
                "currency": pl.Series([row["currency"] for row in records], dtype=pl.String),
                "weight": normalize_weights(weights, declared),
                "shares": pl.Series([row["shares"] for row in records], dtype=pl.Float64),
                "market_value": pl.Series(
                    [row["market_value"] for row in records], dtype=pl.Float64
                ),
                "source_id": pl.Series([raw.source_id] * count, dtype=pl.String),
                "raw_file_sha256": pl.Series([raw.sha256] * count, dtype=pl.String),
            }
        )
        return ParsedFile(
            raw=raw,
            as_of_date=as_of_date,
            published_at=None,
            frame=frame.select(_COLUMNS),
            notes=notes,
        )


PARSERS = [BmoParser()]
