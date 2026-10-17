"""Parsers for the synthetic daily fixture formats.

These parsers read the three daily-file shapes used by the synthetic fixture
set: an iShares-like CSV with a preamble and percent weights, a BMO-like XLSX
with a header block and fraction weights, and a Purpose-like JSON document.
They are also the structural reference for the real issuer parsers: the as-of
date comes from the file itself, the weight unit is declared by the format
(never guessed), and a declared-unit contradiction is reported in the notes.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import polars as pl

from etfh.types import FileFormat, ParsedFile, RawFile
from etfh.units import WeightUnit, check_unit, normalize_weights, parse_number

__all__ = ["PARSERS", "SynCsvParser", "SynJsonParser", "SynXlsxParser"]

_SYN_CSV = "syn_csv"
_SYN_XLSX = "syn_xlsx"
_SYN_JSON = "syn_json"

_COLUMNS = [
    "fund_ticker",
    "as_of_date",
    "security_name",
    "security_ticker",
    "exchange",
    "isin",
    "cusip",
    "asset_class",
    "sector",
    "country",
    "currency",
    "weight",
    "shares",
    "market_value",
    "source_id",
    "raw_file_sha256",
]

_DATE_FORMATS = ("%B %d, %Y", "%b %d, %Y", "%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y")
_HEADER_ALIASES = {
    "security name": "security_name",
    "name": "security_name",
    "ticker": "security_ticker",
    "security ticker": "security_ticker",
    "isin": "isin",
    "cusip": "cusip",
    "asset class": "asset_class",
    "sector": "sector",
    "country": "country",
    "currency": "currency",
    "exchange": "exchange",
    "weight": "weight",
    "weight (%)": "weight",
    "shares": "shares",
    "market value": "market_value",
}


class _SyntheticDailyParser:
    """Shared frame construction for the synthetic daily formats."""

    source_id: str
    formats: frozenset[FileFormat]

    def can_parse(self, raw: RawFile) -> bool:
        return raw.source_id == self.source_id

    def parse(self, raw: RawFile) -> ParsedFile:
        as_of_date, declared, rows = self._read(raw)
        if not rows:
            raise ValueError(f"{raw.path}: no holdings rows found")
        weights = [row["weight"] for row in rows]
        normalized = normalize_weights(weights, declared)
        total = pl.Series("weight", weights, dtype=pl.Float64)
        unit_check = check_unit(total, declared)
        notes = [f"declared_unit={declared.value}"]
        if not unit_check.consistent:
            notes.append(f"unit_check={unit_check.detail()}")
        frame = pl.DataFrame(
            {
                "fund_ticker": pl.Series([raw.fund_ticker] * len(rows), dtype=pl.String),
                "as_of_date": pl.Series([as_of_date] * len(rows), dtype=pl.Date),
                "security_name": pl.Series([row["security_name"] for row in rows], dtype=pl.String),
                "security_ticker": pl.Series(
                    [row["security_ticker"] for row in rows], dtype=pl.String
                ),
                "exchange": pl.Series([row["exchange"] for row in rows], dtype=pl.String),
                "isin": pl.Series([row["isin"] for row in rows], dtype=pl.String),
                "cusip": pl.Series([row["cusip"] for row in rows], dtype=pl.String),
                "asset_class": pl.Series([row["asset_class"] for row in rows], dtype=pl.String),
                "sector": pl.Series([row["sector"] for row in rows], dtype=pl.String),
                "country": pl.Series([row["country"] for row in rows], dtype=pl.String),
                "currency": pl.Series([row["currency"] for row in rows], dtype=pl.String),
                "weight": normalized,
                "shares": pl.Series([row["shares"] for row in rows], dtype=pl.Float64),
                "market_value": pl.Series([row["market_value"] for row in rows], dtype=pl.Float64),
                "source_id": pl.Series([raw.source_id] * len(rows), dtype=pl.String),
                "raw_file_sha256": pl.Series([raw.sha256] * len(rows), dtype=pl.String),
            }
        )
        return ParsedFile(
            raw=raw,
            as_of_date=as_of_date,
            published_at=None,
            frame=frame.select(_COLUMNS),
            notes=notes,
        )

    def _read(self, raw: RawFile) -> tuple[date, WeightUnit, list[dict[str, Any]]]:
        raise NotImplementedError


def parse_date_value(value: Any) -> date | None:
    """Parse a date from a display string, ``date`` or ``datetime`` value."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()  # noqa: DTZ007
        except ValueError:
            continue
    return None


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _clean_identifier(value: Any) -> str | None:
    text = _clean_text(value)
    if text is None:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text.upper()


def _asset_class(value: Any) -> str:
    text = (_clean_text(value) or "").lower()
    if text in {"equity", "cash", "fx", "derivative", "fixed_income", "commodity"}:
        return text
    if text in {"cash and cash equivalents", "cash equivalents"}:
        return "cash"
    if text in {"fixed income", "bond", "bonds"}:
        return "fixed_income"
    if text == "fund" or text == "etf":
        return "other"
    return text or "other"


def _row(record: dict[str, Any]) -> dict[str, Any]:
    weight = parse_number(_clean_text(record.get("weight")))
    if weight is None:
        raise ValueError(f"missing weight for holding {record.get('security_name')!r}")
    shares = parse_number(_clean_text(record.get("shares")))
    market_value = parse_number(_clean_text(record.get("market_value")))
    return {
        "security_name": _clean_text(record.get("security_name")) or "",
        "security_ticker": _clean_identifier(record.get("security_ticker")),
        "exchange": _clean_identifier(record.get("exchange")),
        "isin": _clean_identifier(record.get("isin")),
        "cusip": _clean_identifier(record.get("cusip")),
        "asset_class": _asset_class(record.get("asset_class")),
        "sector": _clean_text(record.get("sector")) or "Unknown",
        "country": _clean_text(record.get("country")) or "Unknown",
        "currency": _clean_identifier(record.get("currency")),
        "weight": weight,
        "shares": shares if shares is not None else 0.0,
        "market_value": market_value,
    }


class SynCsvParser(_SyntheticDailyParser):
    """iShares-like CSV: preamble, then a header row and percent/fraction weights."""

    source_id = _SYN_CSV
    formats = frozenset({FileFormat.csv})

    def _read(self, raw: RawFile) -> tuple[date, WeightUnit, list[dict[str, Any]]]:
        text = Path(raw.path).read_text(encoding="utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        as_of_date = _csv_as_of(rows)
        if as_of_date is None:
            raise ValueError(f"{raw.path}: no 'as of' date found")
        header_index = _csv_header_index(rows)
        headers = [cell.strip() for cell in rows[header_index]]
        declared = (
            WeightUnit.percent if any("%" in cell for cell in headers) else WeightUnit.fraction
        )
        records: list[dict[str, Any]] = []
        for cells in rows[header_index + 1 :]:
            if not any(cell.strip() for cell in cells):
                continue
            record = {
                _HEADER_ALIASES.get(header.lower(), header.lower()): (
                    cells[index] if index < len(cells) else None
                )
                for index, header in enumerate(headers)
            }
            if not _clean_text(record.get("security_name")):
                continue
            if not _clean_text(record.get("weight")):
                continue
            records.append(_row(record))
        return as_of_date, declared, records


def _csv_as_of(rows: list[list[str]]) -> date | None:
    for cells in rows:
        joined = ", ".join(cell.strip() for cell in cells if cell.strip())
        position = joined.lower().find("as of")
        if position < 0:
            continue
        candidate = joined[position + len("as of") :].strip(" \t,;:")
        parsed = parse_date_value(candidate)
        if parsed is not None:
            return parsed
    return None


def _csv_header_index(rows: list[list[str]]) -> int:
    for index, cells in enumerate(rows):
        lowered = {cell.strip().lower() for cell in cells}
        if ("name" in lowered and "weight" in lowered) or "weight (%)" in lowered:
            return index
    raise ValueError("no holdings header row found in CSV")


class SynXlsxParser(_SyntheticDailyParser):
    """BMO-like XLSX: a title block, an 'As of Date' cell, then a header block."""

    source_id = _SYN_XLSX
    formats = frozenset({FileFormat.xlsx})

    def _read(self, raw: RawFile) -> tuple[date, WeightUnit, list[dict[str, Any]]]:
        rows = _xlsx_rows(Path(raw.path))
        as_of_index = next(
            (
                index
                for index, cells in enumerate(rows)
                if any(_clean_text(cell) == "As of Date" for cell in cells)
            ),
            None,
        )
        if as_of_index is None:
            raise ValueError(f"{raw.path}: no 'As of Date' cell found")
        as_of_date = None
        for cell in rows[as_of_index][1:]:
            as_of_date = parse_date_value(cell)
            if as_of_date is not None:
                break
        if as_of_date is None:
            raise ValueError(f"{raw.path}: no as-of date value found")
        header_index = next(
            (
                index
                for index, cells in enumerate(rows)
                if any(_clean_text(cell) == "Security Name" for cell in cells)
            ),
            None,
        )
        if header_index is None:
            raise ValueError(f"{raw.path}: no holdings header row found")
        headers = [_clean_text(cell) or "" for cell in rows[header_index]]
        declared = (
            WeightUnit.percent if any("%" in cell for cell in headers) else WeightUnit.fraction
        )
        records: list[dict[str, Any]] = []
        for cells in rows[header_index + 1 :]:
            record = {
                _HEADER_ALIASES.get(header.lower(), header.lower()): (
                    cells[index] if index < len(cells) else None
                )
                for index, header in enumerate(headers)
            }
            if not _clean_text(record.get("security_name")):
                continue
            if not _clean_text(record.get("weight")):
                continue
            records.append(_row(record))
        return as_of_date, declared, records


def _xlsx_rows(path: Path) -> list[list[Any]]:
    """Read every cell of the first sheet, with an openpyxl fallback."""
    try:
        frame = pl.read_excel(
            path,
            engine="calamine",
            has_header=False,
            infer_schema_length=None,
        )
        return [list(row) for row in frame.rows()]
    except Exception:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        return [list(row) for row in sheet.iter_rows(values_only=True)]


class SynJsonParser(_SyntheticDailyParser):
    """Purpose-like JSON: ``portfolio.dt`` and ``portfolio.holdings[]``."""

    source_id = _SYN_JSON
    formats = frozenset({FileFormat.json})

    def _read(self, raw: RawFile) -> tuple[date, WeightUnit, list[dict[str, Any]]]:
        payload = json.loads(Path(raw.path).read_text(encoding="utf-8"))
        try:
            portfolio = payload["portfolio"]
            holdings = portfolio["holdings"]
        except (KeyError, TypeError) as exc:
            raise ValueError(f"{raw.path}: missing portfolio holdings") from exc
        as_of_date = parse_date_value(portfolio.get("dt"))
        if as_of_date is None:
            raise ValueError(f"{raw.path}: no portfolio date found")
        if not isinstance(holdings, list) or not holdings:
            raise ValueError(f"{raw.path}: portfolio holdings are empty")
        records = [
            _row(
                {
                    "security_name": item.get("name") or item.get("security_name"),
                    "security_ticker": item.get("ticker") or item.get("security_ticker"),
                    "exchange": item.get("exchange"),
                    "isin": item.get("isin"),
                    "cusip": item.get("cusip"),
                    "asset_class": item.get("asset_class"),
                    "sector": item.get("sector"),
                    "country": item.get("country"),
                    "currency": item.get("currency"),
                    "weight": item.get("weight"),
                    "shares": item.get("shares"),
                    "market_value": item.get("market_value"),
                }
            )
            for item in holdings
            if isinstance(item, dict)
        ]
        return as_of_date, WeightUnit.fraction, records


PARSERS: list[SynCsvParser | SynXlsxParser | SynJsonParser] = [
    SynCsvParser(),
    SynXlsxParser(),
    SynJsonParser(),
]
