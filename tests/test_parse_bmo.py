"""Tests for the BMO daily holdings parser."""

from __future__ import annotations

from datetime import date
from pathlib import Path

import polars as pl
import pytest
from openpyxl import Workbook

from etfh.parse.bmo import PARSERS, BmoParser
from etfh.parse.registry import find_parser
from etfh.registry import load_sources
from etfh.schema import ParsedHoldings
from etfh.types import RawFile
from tests._parse_helpers import make_raw

FIXTURE_DIR = Path(__file__).resolve().parent / "fixtures" / "issuers" / "bmo"
CONFIG_DIR = Path(__file__).resolve().parent.parent / "config"


def _raw(name: str, *, fund_ticker: str, as_of: date | None = date(2026, 10, 16)) -> RawFile:
    return make_raw(
        FIXTURE_DIR / name,
        source_id="bmo_daily",
        fund_ticker=fund_ticker,
        as_of_hint=as_of,
    )


def test_parses_equity_workbook() -> None:
    parsed = BmoParser().parse(_raw("equity.xlsx", fund_ticker="EQLK"))
    assert parsed.as_of_date == date(2026, 10, 16)
    assert parsed.published_at is None
    assert parsed.notes[0] == "declared_unit=percent"
    assert "weight_decimals=2" in parsed.notes
    assert not any(note.startswith("unit_check=") for note in parsed.notes)
    frame = parsed.frame
    ParsedHoldings.validate(frame)
    assert frame.height == 5
    assert frame["weight"].to_list() == pytest.approx([0.35, 0.30, 0.20, 0.1494, 0.0006], abs=1e-12)
    assert float(frame["weight"].sum()) == pytest.approx(1.0, abs=1e-9)
    cash = frame.filter(pl.col("asset_class") == "cash")
    assert cash.height == 1
    assert cash["currency"][0] == "CAD"
    assert cash["sector"][0] == "Unknown"
    assert cash["country"][0] == "Unknown"
    assert frame["shares"].null_count() == 0


def test_parses_fixed_income_workbook() -> None:
    parsed = BmoParser().parse(_raw("fixed_income.xlsx", fund_ticker="FILK"))
    frame = parsed.frame
    ParsedHoldings.validate(frame)
    assert frame.height == 4
    assert set(frame["asset_class"].to_list()) == {"fixed_income", "cash"}
    cash = frame.filter(pl.col("asset_class") == "cash")
    assert cash["weight"][0] == pytest.approx(-0.0004, abs=1e-12)
    assert cash["market_value"][0] == pytest.approx(-40000.0)
    assert frame.filter(pl.col("asset_class") == "fixed_income")["shares"].null_count() == 3
    assert float(frame["weight"].sum()) == pytest.approx(1.0, abs=1e-9)
    assert frame["sector"].to_list()[0] == "GOVERNMENT BONDS"


def test_duplicate_isin_lines_are_kept_separate() -> None:
    parsed = BmoParser().parse(_raw("duplicate_lines.xlsx", fund_ticker="DULK"))
    frame = parsed.frame
    ParsedHoldings.validate(frame)
    assert frame.height == 4
    assert float(frame["weight"].sum()) == pytest.approx(1.0, abs=1e-9)
    duplicates = frame.filter(pl.col("isin") == "RU0000000001")
    assert duplicates.height == 2
    assert duplicates["security_name"].unique().to_list() == ["Quarrystone Metals PAO"]
    assert set(duplicates["currency"].to_list()) == {"RUB", "USD"}
    assert duplicates["weight"].to_list() == pytest.approx([0.099, 0.001], abs=1e-12)
    assert duplicates["shares"].to_list() == [5000.0, 1500.0]
    assert duplicates["market_value"].to_list() == pytest.approx([100.0, 50.0])


def test_header_block_variants_are_found_by_name(tmp_path: Path) -> None:
    path = tmp_path / "preamble.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Lookalike Holdings Extract"])
    sheet.append(["Prepared for internal review"])
    sheet.append(["Weight (%)", "Name", "ISIN", "Number of Shares", "Asset Class"])
    sheet.append([100.0, "Northolt Instruments PLC", "CA000000AA00", 10, "Equity"])
    workbook.save(path)
    parsed = BmoParser().parse(
        make_raw(path, source_id="bmo_daily", fund_ticker="EQLK", as_of_hint=date(2026, 10, 16))
    )
    assert parsed.frame.height == 1
    assert parsed.frame["security_name"][0] == "Northolt Instruments PLC"


def test_openpyxl_fallback_parses_the_same_file(monkeypatch: pytest.MonkeyPatch) -> None:
    import etfh.parse.bmo as bmo

    def boom(*args: object, **kwargs: object) -> None:
        raise RuntimeError("calamine unavailable")

    monkeypatch.setattr(bmo.pl, "read_excel", boom)
    parsed = BmoParser().parse(_raw("equity.xlsx", fund_ticker="EQLK"))
    assert parsed.frame.height == 5


def test_unit_contradiction_note_when_values_are_fractions(tmp_path: Path) -> None:
    path = tmp_path / "fractions.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Weight (%)", "Name", "ISIN", "Number of Shares", "Asset Class"])
    sheet.append([0.6, "Northolt Instruments PLC", "CA000000AA00", 10, "Equity"])
    sheet.append([0.4, "Pinehurst Utilities Inc.", "CA000000BB11", 20, "Equity"])
    workbook.save(path)
    parsed = BmoParser().parse(
        make_raw(path, source_id="bmo_daily", fund_ticker="EQLK", as_of_hint=date(2026, 10, 16))
    )
    notes = [note for note in parsed.notes if note.startswith("unit_check=")]
    assert len(notes) == 1
    assert "total matches unit fraction" in notes[0]


def test_weight_decimals_note_tracks_the_printed_precision(tmp_path: Path) -> None:
    path = tmp_path / "precision.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Weight (%)", "Name", "ISIN", "Number of Shares", "Asset Class"])
    sheet.append([60.125, "Northolt Instruments PLC", "CA600001AA14", 10, "Equity"])
    sheet.append([39.875, "Pinehurst Utilities Inc.", "CA600002AA96", 20, "Equity"])
    workbook.save(path)
    parsed = BmoParser().parse(
        make_raw(path, source_id="bmo_daily", fund_ticker="EQLK", as_of_hint=date(2026, 10, 16))
    )
    assert "weight_decimals=3" in parsed.notes


def test_missing_as_of_hint_raises() -> None:
    with pytest.raises(ValueError, match="no as-of date"):
        BmoParser().parse(_raw("equity.xlsx", fund_ticker="EQLK", as_of=None))


def test_discovered_by_the_parser_registry() -> None:
    source = load_sources(CONFIG_DIR / "sources.yaml")["bmo_daily"]
    parser = find_parser(_raw("equity.xlsx", fund_ticker="EQLK"), source)
    assert isinstance(parser, BmoParser)


def test_can_parse_only_bmo_xlsx() -> None:
    parser = BmoParser()
    assert parser.can_parse(_raw("equity.xlsx", fund_ticker="EQLK"))
    other = make_raw(FIXTURE_DIR / "equity.xlsx", source_id="syn_xlsx", fund_ticker="EQLK")
    assert not parser.can_parse(other)
    assert len(PARSERS) == 1
    assert isinstance(PARSERS[0], BmoParser)
