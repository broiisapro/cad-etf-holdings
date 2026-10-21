"""Test-only readers that build ``ParsedFile``s from the committed fixtures.

The real parsers live in ``etfh/parse`` (workstreams A and B); this module
mirrors the interfaces.md conventions (notes, declared units, security ids)
closely enough to exercise every validation check without depending on parser
code that may not be merged yet. Fixture numbers are never re-derived here:
every weight, name and identifier comes from the stored fixture file.
"""

from __future__ import annotations

import csv
import json
from datetime import UTC, date, datetime, time
from pathlib import Path
from time import strptime
from typing import Any
from zoneinfo import ZoneInfo

import pdfplumber
import polars as pl
from openpyxl import load_workbook

from etfh.storage import RawStore
from etfh.types import ParsedFile, RawFile
from etfh.units import WeightUnit

FIXTURES = Path(__file__).resolve().parent / "fixtures"
RAW_ROOT = FIXTURES / "raw_synthetic"
MANIFEST_PATH = FIXTURES / "manifest.json"
_TORONTO = ZoneInfo("America/Toronto")
_STRING_COLUMNS = ("security_ticker", "exchange", "isin", "cusip", "currency")


def load_manifest() -> dict[str, Any]:
    data = json.loads(MANIFEST_PATH.read_text(encoding="utf-8"))
    assert isinstance(data, dict)
    return data


def raw_files() -> list[RawFile]:
    return RawStore(RAW_ROOT).iter_raw()


def _isin_check_digit(payload: str) -> str:
    digits = "".join(str(int(character, 36)) for character in payload.upper())
    total = 0
    for index, character in enumerate(reversed(digits)):
        value = int(character)
        if index % 2 == 0:
            value *= 2
            if value > 9:
                value -= 9
        total += value
    return str((10 - total % 10) % 10)


def _cusip_to_isin(cusip: str, country: str) -> str | None:
    if country not in {"CA", "US"} or len(cusip) != 9:
        return None
    payload = country + cusip
    return payload + _isin_check_digit(payload)


def _security_id(row: dict[str, object]) -> str | None:
    asset_class = str(row["asset_class"])
    if asset_class in {"cash", "fx", "derivative"}:
        currency = row["currency"]
        if asset_class == "cash" and currency:
            return f"CASH:{currency}"
        return f"DERIV:{row['fund_ticker']}:{row['security_name']}"
    isin = row["isin"]
    if isin:
        return f"ISIN:{isin}"
    cusip = row["cusip"]
    if cusip:
        derived = _cusip_to_isin(str(cusip), str(row["country"]))
        return f"ISIN:{derived}" if derived else f"CUSIP:{cusip}"
    ticker = row["security_ticker"]
    exchange = row["exchange"]
    if ticker and exchange:
        return f"TKR:{exchange}:{ticker}"
    return None


def _frame(rows: list[dict[str, object]]) -> pl.DataFrame:
    overrides: dict[str, pl.DataType] = {column: pl.String for column in _STRING_COLUMNS}
    overrides.update({"weight": pl.Float64, "shares": pl.Float64, "market_value": pl.Float64})
    frame = pl.DataFrame(rows, schema_overrides=overrides)
    security_ids = [_security_id(row) for row in rows]
    return frame.with_columns(pl.Series("security_id", security_ids, dtype=pl.String))


def raw_decimals(value: object) -> int:
    """Decimals printed for one raw weight value (``7.77`` -> 2, ``0.0908`` -> 4)."""
    text = repr(value) if isinstance(value, float) else str(value)
    text = text.strip().rstrip("%")
    if "." not in text:
        return 0
    return len(text.split(".", 1)[1])


def printed_decimals(values: list[object]) -> int:
    """The most decimals printed among a file's raw weight values."""
    return max((raw_decimals(value) for value in values), default=0)


def _base_row(raw: RawFile, as_of: date, **values: object) -> dict[str, object]:
    row: dict[str, object] = {
        "fund_ticker": raw.fund_ticker,
        "as_of_date": as_of,
        "security_name": "",
        "security_ticker": None,
        "exchange": None,
        "isin": None,
        "cusip": None,
        "asset_class": "equity",
        "sector": "Unknown",
        "country": "Unknown",
        "currency": None,
        "weight": 0.0,
        "shares": 0.0,
        "market_value": None,
        "source_id": raw.source_id,
        "raw_file_sha256": raw.sha256,
    }
    row.update(values)
    return row


def _parse_csv(raw: RawFile) -> ParsedFile:
    text = raw.path.read_text(encoding="utf-8")
    lines = text.splitlines()
    as_of: date | None = None
    header_index: int | None = None
    for index, row in enumerate(csv.reader(lines)):
        if row and row[0] == "Fund Holdings as of" and len(row) > 1:
            parsed_date = strptime(row[1], "%B %d, %Y")
            as_of = date(parsed_date.tm_year, parsed_date.tm_mon, parsed_date.tm_mday)
        if row[:2] == ["Ticker", "Name"]:
            header_index = index
    assert as_of is not None, f"{raw.path}: no as-of date"
    assert header_index is not None, f"{raw.path}: no holdings header"
    reader = csv.DictReader(lines[header_index:])
    unit = WeightUnit.percent if "Weight (%)" in (reader.fieldnames or []) else WeightUnit.fraction
    weight_column = "Weight (%)" if unit is WeightUnit.percent else "Weight"
    parsed_rows: list[dict[str, object]] = []
    weight_texts: list[object] = []
    for row in reader:
        weight_texts.append(row[weight_column])
        weight = float(row[weight_column]) / (100.0 if unit is WeightUnit.percent else 1.0)
        parsed_rows.append(
            _base_row(
                raw,
                as_of,
                security_name=row["Name"],
                security_ticker=row["Ticker"] or None,
                exchange=row["Exchange"] or None,
                isin=row["ISIN"] or None,
                asset_class=row["Asset Class"].lower(),
                sector=row["Sector"] or "Unknown",
                country=row["Country"] or "Unknown",
                currency=row["Currency"] or None,
                weight=weight,
                shares=float(row["Shares"]),
                market_value=float(row["Market Value"]),
            )
        )
    return ParsedFile(
        raw=raw,
        as_of_date=as_of,
        published_at=None,
        frame=_frame(parsed_rows),
        notes=[
            f"declared_unit={unit.value}",
            f"weight_decimals={printed_decimals(weight_texts)}",
        ],
    )


def _parse_xlsx(raw: RawFile) -> ParsedFile:
    workbook = load_workbook(raw.path, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    as_of: date | None = None
    header_index: int | None = None
    for index, row in enumerate(rows):
        if row[:1] == ("As of Date",) and row[1] is not None:
            as_of = date.fromisoformat(str(row[1]))
        if row[:1] == ("Security Name",):
            header_index = index
    assert as_of is not None, f"{raw.path}: no as-of date"
    assert header_index is not None, f"{raw.path}: no holdings header"
    parsed_rows: list[dict[str, object]] = []
    weight_values: list[object] = []
    for row in rows[header_index + 1 :]:
        if row[0] is None:
            continue
        weight_values.append(row[8])
        parsed_rows.append(
            _base_row(
                raw,
                as_of,
                security_name=str(row[0]),
                security_ticker=row[1],
                isin=row[2],
                cusip=row[3],
                exchange=row[4],
                sector=row[5] or "Unknown",
                country=row[6] or "Unknown",
                currency=row[7],
                weight=float(row[8]),
                shares=float(row[9]),
                market_value=float(row[10]),
            )
        )
    return ParsedFile(
        raw=raw,
        as_of_date=as_of,
        published_at=None,
        frame=_frame(parsed_rows),
        notes=[
            "declared_unit=fraction",
            f"weight_decimals={printed_decimals(weight_values)}",
        ],
    )


def _parse_json(raw: RawFile) -> ParsedFile:
    payload = json.loads(raw.path.read_text(encoding="utf-8"))
    as_of = date.fromisoformat(payload["portfolio"]["dt"])
    holdings = payload["portfolio"]["holdings"]
    parsed_rows = [
        _base_row(
            raw,
            as_of,
            security_name=holding["name"],
            security_ticker=holding["ticker"],
            isin=holding["isin"],
            cusip=holding["cusip"],
            exchange=holding["exchange"],
            asset_class=holding["asset_class"],
            sector=holding["sector"] or "Unknown",
            country=holding["country"] or "Unknown",
            currency=holding["currency"],
            weight=float(holding["weight"]),
            shares=float(holding["shares"]),
            market_value=float(holding["market_value"]),
        )
        for holding in holdings
    ]
    return ParsedFile(
        raw=raw,
        as_of_date=as_of,
        published_at=None,
        frame=_frame(parsed_rows),
        notes=[
            "declared_unit=fraction",
            f"weight_decimals={printed_decimals([holding['weight'] for holding in holdings])}",
        ],
    )


def _parse_pdf(raw: RawFile) -> ParsedFile:
    with pdfplumber.open(raw.path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    as_of: date | None = None
    published_at: datetime | None = None
    stated_total: float | None = None
    table: list[tuple[str, float]] = []
    weight_texts: list[object] = []
    in_table = False
    for line in text.splitlines():
        if line.startswith("Summary of Investment Portfolio as at "):
            as_of = date.fromisoformat(line.rsplit(" ", 1)[1])
        elif line.startswith("Published "):
            published_on = date.fromisoformat(line.split(" ", 1)[1])
            published_at = datetime.combine(published_on, time(0, 0), tzinfo=_TORONTO).astimezone(
                UTC
            )
        elif line == "Holding % of net asset value":
            in_table = True
        elif line.startswith("Total of the holdings listed above:"):
            stated_total = float(line.split(":")[1].strip().split("%")[0]) / 100.0
        elif in_table and line != "-----":
            name, _, value_text = line.rpartition(" ")
            try:
                value = float(value_text)
            except ValueError:
                continue
            weight_texts.append(value_text)
            table.append((name.strip(), value / 100.0))
    assert as_of is not None, f"{raw.path}: no as-of date"
    assert stated_total is not None, f"{raw.path}: no stated total"
    parsed_rows = [
        _base_row(
            raw,
            as_of,
            security_name=security_name,
            asset_class="cash" if security_name == "Cash and Cash Equivalents" else "equity",
            currency="CAD" if security_name == "Cash and Cash Equivalents" else None,
            weight=weight,
        )
        for security_name, weight in table
    ]
    return ParsedFile(
        raw=raw,
        as_of_date=as_of,
        published_at=published_at,
        frame=_frame(parsed_rows),
        notes=[
            "declared_unit=percent",
            f"weight_decimals={printed_decimals(weight_texts)}",
            f"stated_total_fraction={stated_total}",
            "partial_list=true",
        ],
    )


def parse_fixture(raw: RawFile) -> ParsedFile:
    if raw.ext == "csv":
        return _parse_csv(raw)
    if raw.ext == "xlsx":
        return _parse_xlsx(raw)
    if raw.ext == "json":
        return _parse_json(raw)
    if raw.ext == "pdf":
        return _parse_pdf(raw)
    raise ValueError(f"{raw.path}: unsupported fixture format {raw.ext!r}")
